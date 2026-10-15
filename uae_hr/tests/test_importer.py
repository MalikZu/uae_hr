import os
import tempfile

import frappe

from uae_hr import importer
from uae_hr.tests.utils import EID_1990, IBAN_OK, MAINLAND, UAETestCase, make_company, make_employee


class TestImporter(UAETestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		make_company(MAINLAND)

	def setUp(self):
		super().setUp()
		self.dir = tempfile.mkdtemp()
		self.path = os.path.join(self.dir, "import.xlsx")
		importer.make_template(self.path)
		self.emp = make_employee(date_of_birth="1990-05-01")

	def fill(self, sheet, rows):
		from openpyxl import load_workbook

		wb = load_workbook(self.path)
		ws = wb[sheet]
		labels = [c.value for c in ws[1]]
		fields = {f: labels.index(label) for f, label in importer.SHEETS[sheet]["columns"]}
		for row in rows:
			values = [None] * len(labels)
			for k, v in row.items():
				values[fields[k]] = v
			ws.append(values)
		wb.save(self.path)

	def messages(self, result, level=None):
		return [(i["sheet"], i["field"]) for i in result["items"] if level in (None, i["level"])]

	def test_template_has_every_sheet(self):
		from openpyxl import load_workbook

		self.assertEqual(
			load_workbook(self.path).sheetnames,
			["README", "Employees", "Documents", "Issued Items", "Salary"],
		)

	def test_clean_workbook_has_no_errors(self):
		self.fill(
			"Employees",
			[
				{
					"name": self.emp.name,
					"uae_nationality": "India",
					"uae_probation_months": 6,
					"notice_number_of_days": 30,
					"iban": IBAN_OK,
					"uae_person_code": "12345678901234",
					"salary_mode": "Bank",
				}
			],
		)
		self.fill(
			"Documents",
			[
				{
					"employee": self.emp.name,
					"document_type": "Passport",
					"document_number": "Z111",
					"issue_date": "2022-01-01",
					"expiry_date": "2032-01-01",
				},
				{
					"employee": self.emp.name,
					"document_type": "Emirates ID",
					"document_number": EID_1990,
					"expiry_date": "2030-01-01",
				},
				{
					"employee": self.emp.name,
					"document_type": "Residence Visa",
					"passport_number": "Z111",
					"expiry_date": "2030-01-01",
				},
			],
		)
		self.fill("Salary", [{"employee": self.emp.name, "uae_basic": 5000, "uae_housing": 2000}])
		result = importer.check(self.path, write_report=False)
		self.assertEqual(result["errors"], 0, result["items"])

	def test_errors_are_found(self):
		self.fill(
			"Employees",
			[
				{
					"name": self.emp.name,
					"uae_probation_months": 8,
					"notice_number_of_days": 7,
					"uae_nationality": "Atlantis",
					"uae_death_nominee": "X",
					"uae_death_nominee_relation": "Cousin",
				}
			],
		)
		self.fill(
			"Documents",
			[
				{
					"employee": "NO-SUCH-EMP",
					"document_type": "Passport",
					"document_number": "A1",
					"expiry_date": "2030-01-01",
				},
				{"employee": self.emp.name, "document_type": "Passport", "document_number": "A2"},
				{
					"employee": self.emp.name,
					"document_type": "Emirates ID",
					"document_number": "784-1985-1234567-1",
					"expiry_date": "2030-01-01",
				},
				{
					"employee": self.emp.name,
					"document_type": "Residence Visa",
					"passport_number": "OTHER",
					"expiry_date": "2030-01-01",
				},
			],
		)
		self.fill("Issued Items", [{"name": self.emp.name, "item_type": "Laptop"}])
		self.fill("Salary", [{"employee": self.emp.name, "uae_housing": 2000}])
		result = importer.check(self.path)
		errors = self.messages(result, "error")
		for expected in (
			("Employees", "uae_probation_months"),
			("Employees", "notice_number_of_days"),
			("Employees", "uae_nationality"),
			("Employees", "uae_death_nominee_relation"),
			("Documents", "employee"),
			("Documents", "expiry_date"),
			("Salary", "uae_basic"),
		):
			self.assertIn(expected, errors)
		warns = self.messages(result, "warning")
		self.assertIn(("Documents", "document_number"), warns)
		self.assertIn(("Documents", "passport_number"), warns)
		self.assertTrue(os.path.exists(result["report"]))

	def test_split_for_data_import(self):
		self.fill("Employees", [{"name": self.emp.name, "uae_probation_months": 3}])
		self.fill("Documents", [{"employee": self.emp.name, "document_type": "NDA", "document_number": "N1"}])
		written = {w["sheet"]: w for w in importer.split(self.path, os.path.join(self.dir, "out"))}
		self.assertEqual(set(written), {"Employees", "Documents"})
		self.assertEqual(written["Employees"]["import_type"], "Update Existing Records")
		from openpyxl import load_workbook

		header = [c.value for c in load_workbook(written["Employees"]["file"]).active[1]]
		self.assertEqual(header, ["ID", "Probation (months)"])

	def test_documents_sheet_validates_in_data_import(self):
		"""The split Documents file is accepted by standard Data Import (allow_import was off in 0.1.0)."""
		self.fill(
			"Documents",
			[{"employee": self.emp.name, "document_type": "NDA", "document_number": "N1", "notes": "Signed"}],
		)
		docs = {w["sheet"]: w for w in importer.split(self.path, os.path.join(self.dir, "out"))}["Documents"]
		with open(docs["file"], "rb") as fh:
			file_doc = frappe.get_doc(
				{"doctype": "File", "file_name": "uae_documents.xlsx", "content": fh.read(), "is_private": 1}
			).insert(ignore_permissions=True)
		# the rollback removes the File row, not the file on disk
		on_disk = file_doc.get_full_path()
		self.addCleanup(lambda: os.path.exists(on_disk) and os.remove(on_disk))

		user = "uae.import.test@example.com"
		if not frappe.db.exists("User", user):
			frappe.get_doc(
				{
					"doctype": "User",
					"email": user,
					"first_name": "Import",
					"roles": [{"role": "System Manager"}],
				}
			).insert(ignore_permissions=True)
		with self.set_user(user):
			di = frappe.get_doc(
				{
					"doctype": "Data Import",
					"reference_doctype": docs["doctype"],
					"import_type": docs["import_type"],
					"import_file": file_doc.file_url,
				}
			).insert()
			preview = di.get_preview_from_template()
		self.assertEqual(di.payload_count, 1)
		self.assertEqual([w for w in preview.warnings if w.get("type") != "info"], [])
		mapped = [c.df["fieldname"] for c in preview.columns if c.get("df") and not c.skip_import]
		self.assertEqual(sorted(mapped), sorted(["employee", "document_type", "document_number", "notes"]))


class TestImporterRenewal(UAETestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		make_company(MAINLAND)

	def test_renewal_rows_are_not_duplicates(self):
		d = tempfile.mkdtemp()
		path = os.path.join(d, "import.xlsx")
		importer.make_template(path)
		emp = make_employee()
		TestImporter.fill(
			frappe._dict(path=path),
			"Documents",
			[
				{
					"employee": emp.name,
					"document_type": "Emirates ID",
					"document_number": EID_1990,
					"expiry_date": "2027-01-01",
				},
				{
					"employee": emp.name,
					"document_type": "Emirates ID",
					"document_number": EID_1990,
					"expiry_date": "2030-01-01",
				},
				{
					"employee": emp.name,
					"document_type": "Emirates ID",
					"document_number": EID_1990,
					"expiry_date": "2030-01-01",
				},
			],
		)
		result = importer.check(path, write_report=False)
		dups = [
			i["row"] for i in result["items"] if i["field"] == "document_number" and i["level"] == "error"
		]
		self.assertEqual(dups, [4])


class TestImporterNewFields(UAETestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		make_company(MAINLAND)

	def setUp(self):
		super().setUp()
		d = tempfile.mkdtemp()
		self.path = os.path.join(d, "import.xlsx")
		importer.make_template(self.path)
		self.emp = make_employee()

	def fill(self, sheet, rows):
		TestImporter.fill(self, sheet, rows)

	def test_template_has_the_new_columns(self):
		from openpyxl import load_workbook

		wb = load_workbook(self.path)
		emp_cols = [c.value for c in wb["Employees"][1]]
		for label in ("Visa Establishment", "Visa Emirate", "Engagement Type", "Payment Method"):
			self.assertIn(label, emp_cols)
		doc_cols = [c.value for c in wb["Documents"][1]]
		for label in ("Basic Wage", "Housing Allowance", "Transport Allowance", "Other Allowance"):
			self.assertIn(label, doc_cols)

	def test_new_fields_are_checked(self):
		self.fill(
			"Employees",
			[
				{
					"name": self.emp.name,
					"uae_engagement_type": "Salaried",
					"uae_visa_emirate": "Atlantis",
					"uae_visa_establishment": "No Such Sponsor",
					"uae_payment_method": "Barter",
				}
			],
		)
		self.fill(
			"Documents",
			[
				{
					"employee": self.emp.name,
					"document_type": "MOHRE Employment Contract",
					"document_number": "C1",
					"expiry_date": "2028-01-01",
					"contract_basic": 5000,
					"contract_housing": -10,
				}
			],
		)
		self.fill("Salary", [{"employee": self.emp.name, "uae_basic": 4000}])
		result = importer.check(self.path, write_report=False)
		errors = {(i["sheet"], i["field"]) for i in result["items"] if i["level"] == "error"}
		for expected in (
			("Employees", "uae_visa_emirate"),
			("Employees", "uae_visa_establishment"),
			("Employees", "uae_payment_method"),
			("Documents", "contract_housing"),
		):
			self.assertIn(expected, errors)
		warns = {(i["sheet"], i["field"]) for i in result["items"] if i["level"] == "warning"}
		self.assertIn(("Salary", "uae_basic"), warns)  # below the contract basic

	def test_salaried_without_payment_method_warns(self):
		self.fill("Employees", [{"name": self.emp.name, "uae_engagement_type": "Salaried"}])
		result = importer.check(self.path, write_report=False)
		self.assertIn(
			("Employees", "uae_payment_method", "warning"),
			{(i["sheet"], i["field"], i["level"]) for i in result["items"]},
		)
